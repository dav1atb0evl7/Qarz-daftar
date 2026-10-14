import io
import csv
from decimal import Decimal
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from django.http import HttpResponse
from django.utils import timezone


class ReportExporter:
    @staticmethod
    def export_csv(filename, headers, rows):
        """Generates a downloadable CSV HTTP response with UTF-8 BOM."""
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
        response.write('\ufeff')  # Excel UTF-8 BOM

        writer = csv.writer(response)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(row)

        return response

    @staticmethod
    def export_excel(filename, sheet_title, headers, rows, summary_metrics=None):
        """Generates a styled Excel workbook HTTP response using openpyxl."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_title[:31]

        # Colors & Fonts
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")

        current_row = 1

        # Title
        ws.cell(row=current_row, column=1, value=filename.replace('_', ' ').upper()).font = title_font
        current_row += 1
        ws.cell(row=current_row, column=1, value=f"Generated: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}")
        current_row += 2

        # Optional summary cards
        if summary_metrics:
            for k, v in summary_metrics.items():
                ws.cell(row=current_row, column=1, value=str(k)).font = Font(bold=True)
                ws.cell(row=current_row, column=2, value=str(v))
                current_row += 1
            current_row += 1

        # Table Headers
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        current_row += 1

        # Rows
        thin_border = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB')
        )

        for row_data in rows:
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=str(val) if isinstance(val, Decimal) else val)
                cell.border = thin_border
                if isinstance(val, (int, float, Decimal)):
                    cell.alignment = Alignment(horizontal="right")
            current_row += 1

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
        return response

    @staticmethod
    def export_pdf(filename, title, headers, rows, summary_metrics=None):
        """Generates a downloadable PDF report using reportlab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1E3A8A'),
            spaceAfter=6
        )
        sub_style = ParagraphStyle(
            'ReportSubtitle',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#6B7280'),
            spaceAfter=15
        )

        elements = []
        elements.append(Paragraph(title, title_style))
        elements.append(Paragraph(f"Chiqarilgan sana: {timezone.now().strftime('%Y-%m-%d %H:%M:%S')}", sub_style))

        if summary_metrics:
            summary_data = [[Paragraph(f"<b>{k}</b>", styles['Normal']), str(v)] for k, v in summary_metrics.items()]
            sum_table = Table(summary_data, colWidths=[200, 300])
            sum_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F3F4F6')),
                ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
            ]))
            elements.append(sum_table)
            elements.append(Spacer(1, 15))

        # Main Data Table
        table_data = [headers]
        for row in rows:
            clean_row = [str(cell) for cell in row]
            table_data.append(clean_row)

        col_width = (A4[0] - 60) / max(len(headers), 1)
        data_table = Table(table_data, colWidths=[col_width] * len(headers), repeatRows=1)
        data_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')])
        ]))

        elements.append(data_table)
        doc.build(elements)
        buffer.seek(0)

        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}.pdf"'
        return response
